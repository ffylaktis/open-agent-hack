"""Build the submission slide deck (PPTX) for the Evidence-Intelligence Agent.

    pip install python-pptx
    python deck/build_deck.py        ->  deck/Evidence-Intelligence-Agent.pptx

Content mirrors docs/slides-outline.md; the numbers are from the recorded run
(data/sample_run.json). Regenerate any time the story changes — it's deterministic.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK = RGBColor(0x14, 0x20, 0x2B)
TEAL = RGBColor(0x12, 0x8A, 0x98)
GOLD = RGBColor(0xB8, 0x84, 0x2A)
RED = RGBColor(0xC0, 0x39, 0x2B)
RED_BRIGHT = RGBColor(0xE2, 0x5C, 0x4D)   # reads on the dark slides
MUTED = RGBColor(0x5B, 0x6B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE3, 0xE9, 0xEF)
DIM = RGBColor(0xA9, 0xB7, 0xC0)
FONT = "Segoe UI"
FOOT = ("Evidence-Intelligence Agent · Open Agent Hackathon 2026 · Track 02 · "
        "Fylaktis Fylaktou · UCLA Semel Institute")
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
UI_FULL = os.path.join(ASSETS, "ui_full.png")             # captured by deck/capture_ui.py
UI_REJECTED = os.path.join(ASSETS, "ui_rejected.png")
_LOGO = os.path.join(ASSETS, "logos")
L_NVIDIA = os.path.join(_LOGO, "nvidia.png")
L_ZETARIS = os.path.join(_LOGO, "zetaris.png")
L_METERLESS = os.path.join(_LOGO, "meterless.png")
L_CURSOR = os.path.join(_LOGO, "cursor.png")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _run(p, text, size, color, bold=False, italic=False, strike=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT
    if strike:
        try:
            r.font._rPr.set("strike", "sngStrike")
        except Exception:
            pass
    return r


def _box(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tb.text_frame.word_wrap = True
    tb.text_frame.vertical_anchor = anchor
    return tb.text_frame


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _rect(slide, left, top, width, height, color, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _conn(slide, x1, y1, x2, y2, color, width_pt=2.0, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width_pt)
    if arrow:
        try:
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        except Exception:
            pass
    return c


def _footer(slide):
    _rect(slide, 0.7, 7.0, 11.93, Pt(1.2) / Inches(1), LINE)
    tf = _box(slide, 0.7, 7.04, 11.93, 0.35)
    _run(tf.paragraphs[0], FOOT, 9, MUTED)


def _header(slide, kicker, title, accent=TEAL):
    tf = _box(slide, 0.7, 0.5, 11.9, 0.4)
    _run(tf.paragraphs[0], kicker.upper(), 13, accent, bold=True)
    tf = _box(slide, 0.7, 0.95, 12.0, 1.0)
    _run(tf.paragraphs[0], title, 30, INK, bold=True)
    _rect(slide, 0.72, 1.92, 1.1, 0.055, accent)


def _bullets(slide, bullets, accent=TEAL, top=2.3, size=18, gap=12, left=0.7, width=11.9,
             anchor=MSO_ANCHOR.TOP, height=4.5):
    tf = _box(slide, left, top, width, height, anchor=anchor)
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()           # hanging indent: wrapped lines align under the text, not the ▸
        pPr.set("marL", str(int(Inches(0.33))))
        pPr.set("indent", str(int(Inches(-0.33))))
        _run(p, "▸  ", size, accent, bold=True)
        if isinstance(b, tuple):
            _run(p, b[0], size, INK, bold=True)
            if b[1]:
                _run(p, " " + b[1], size, MUTED)
        else:
            _run(p, b, size, INK)


def content_slide(kicker, title, bullets, accent=TEAL):
    s = prs.slides.add_slide(BLANK)
    _bg(s, WHITE)
    _header(s, kicker, title, accent)
    _bullets(s, bullets, accent, top=2.15, height=4.7, anchor=MSO_ANCHOR.MIDDLE)
    _footer(s)
    return s


def _picture(slide, path, left, top, max_w, max_h):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    w, h = max_w, max_w / ar
    if h > max_h:
        h, w = max_h, max_h * ar
    pic = slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(w), height=Inches(h))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(1)
    return pic


def _stamp(slide, left, top, w, h, text, color, rot=-8, size=15):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    sp.fill.background()
    sp.line.color.rgb = color
    sp.line.width = Pt(2.25)
    sp.rotation = rot
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, text, size, color, bold=True)
    return sp


def _logo_cards(slide, paths, top, cell_h, left=0.9, total_w=11.53):
    from PIL import Image
    n = len(paths)
    cell_w = total_w / n
    pad = 0.2
    for i, path in enumerate(paths):
        cx0 = left + i * cell_w
        card = _rect(slide, cx0 + 0.12, top, cell_w - 0.24, cell_h, RGBColor(0xFA, 0xFB, 0xFC),
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        card.line.color.rgb = LINE
        card.line.width = Pt(1)
        if not os.path.exists(path):
            continue
        with Image.open(path) as im:
            iw, ih = im.size
        ar = iw / ih
        avail_w, avail_h = cell_w - 0.24 - 2 * pad, cell_h - 2 * pad
        w, h = avail_w, avail_w / ar
        if h > avail_h:
            h, w = avail_h, avail_h * ar
        ccx, ccy = cx0 + cell_w / 2, top + cell_h / 2
        slide.shapes.add_picture(path, Inches(ccx - w / 2), Inches(ccy - h / 2), width=Inches(w), height=Inches(h))


# 1 — Title
s = prs.slides.add_slide(BLANK)
_bg(s, INK)
tf = _box(s, 0.9, 0.8, 11.5, 0.5)
_run(tf.paragraphs[0], "OPEN AGENT HACKATHON 2026 · TRACK 02", 14, TEAL, bold=True)
_rect(s, 0.92, 2.45, 1.4, 0.06, TEAL)
tf = _box(s, 0.9, 2.7, 11.6, 2.2)
_run(tf.paragraphs[0], "The Agent That Can Explain Why", 46, WHITE, bold=True)
p = tf.add_paragraph()
p.space_before = Pt(14)
_run(p, "An evidence-intelligence agent: it grounds every clinical claim in real data — "
        "and deletes the ones it can't prove.", 20, DIM)
tf = _box(s, 0.9, 6.35, 11.6, 0.6)
_run(tf.paragraphs[0], "Fylaktis Fylaktou · UCLA Semel Institute      Zetaris · Meterless · NVIDIA Nemotron · Cursor",
     13, DIM)

# 2 — Problem & users
content_slide("The problem · Impact", "Confident, uncited, and wrong", [
    ("LLMs fabricate clinical citations.", "Published studies report 19.9–91% inaccurate or invented references."),
    ("An unverifiable answer is worse than none.", "In medicine and regulation you must know what's grounded."),
    ("Evidence synthesis is slow.", "Reconciling trial registries against device clearances is manual and costly."),
    ("Users:", "clinical & regulatory evidence teams, pharma competitive-intelligence, researchers — who need an "
              "auditable answer, not a chatbot guess."),
])

# 3 — The demo question
s = prs.slides.add_slide(BLANK)
_bg(s, WHITE)
tf = _box(s, 0.7, 0.5, 11.9, 0.4)
_run(tf.paragraphs[0], "THE DEMO QUESTION", 13, TEAL, bold=True)
_rect(s, 0.72, 1.0, 1.1, 0.055, TEAL)
tf = _box(s, 0.9, 1.95, 11.5, 2.0, anchor=MSO_ANCHOR.MIDDLE)
_run(tf.paragraphs[0],
     "“What neuromodulation is being trialed for anorexia nervosa, "
     "and is any of it FDA-cleared for that use?”", 32, INK, bold=True, italic=True)
tf = _box(s, 0.9, 4.25, 11.5, 1.0)
_run(tf.paragraphs[0],
     "A real cross-source question where the honest answer is nuanced — and easy to get confidently wrong.",
     18, MUTED)
_footer(s)

# 4 — Solution (text left; hero UI screenshot right when captured, else full-width text)
S4_BULLETS = [
    ("Grounded, not guessed.", "Gathers governed evidence and cites the exact SQL behind every claim."),
    ("An Evaluator that deletes what it can't back.", "Each claim is re-checked against only its own evidence."),
    ("Honest scope.", "Affirms the scoped conclusion; rejects the over-claim "
                      "“rTMS is FDA-cleared for anorexia.”"),
]
if os.path.exists(UI_FULL):
    s = prs.slides.add_slide(BLANK)
    _bg(s, WHITE)
    _header(s, "The solution", "Grounded answers, with receipts")
    _bullets(s, S4_BULLETS, top=2.35, left=0.7, width=5.95, size=17, gap=16)
    _picture(s, UI_FULL, left=7.05, top=2.2, max_w=5.6, max_h=4.55)
    _footer(s)
else:
    content_slide("The solution", "Grounded answers, with receipts", S4_BULLETS)

# 5 — Architecture
s = prs.slides.add_slide(BLANK)
_bg(s, WHITE)
_header(s, "Architecture · Technical", "A real agent, not a pipeline")
labels = ["Planner", "Analyst", "cross_source", "Reporter", "Evaluator"]
subs = ["decompose", "SQL → evidence", "trialed × cleared", "cite claims", "verify / reject"]
n, margin, gap, box_h, y = 5, 0.7, 0.42, 1.15, 2.5
total_w = 13.333 - 2 * margin
box_w = (total_w - (n - 1) * gap) / n
centers = []
for i, (lab, sub) in enumerate(zip(labels, subs, strict=True)):
    x = margin + i * (box_w + gap)
    centers.append(x + box_w / 2)
    hot = lab == "Evaluator"
    sp = _rect(s, x, y, box_w, box_h, TEAL if hot else WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    if not hot:
        sp.line.color.rgb = TEAL
        sp.line.width = Pt(1.5)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, lab, 14, WHITE if hot else INK, bold=True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    _run(p2, sub, 9.5, RGBColor(0xDD, 0xF0, 0xF2) if hot else MUTED)
    if i < n - 1:
        _rect(s, x + box_w + 0.02, y + box_h / 2 - 0.1, gap - 0.04, 0.2, TEAL, shape=MSO_SHAPE.RIGHT_ARROW)

# the feedback loop (the differentiator) drawn as a REAL return edge: Evaluator -> back to Reporter
bot, lane = y + box_h, y + box_h + 0.7
ev_c, rp_c = centers[4], centers[3]
_conn(s, ev_c, bot, ev_c, lane, GOLD, arrow=False)
_conn(s, ev_c, lane, rp_c, lane, GOLD, arrow=False)
_conn(s, rp_c, lane, rp_c, bot, GOLD, arrow=True)
loop_lines = [
    (lane + 0.12, GOLD,
     "✗ unsupported → the Reporter revises & re-cites, then re-evaluates (loops ≤ MAX_ROUNDS)"),
    (lane + 0.52, TEAL,
     "✓ the cited answer ships only when every claim passes the Evaluator"),
]
for ty, col, txt in loop_lines:
    t = _box(s, 0.7, ty, 11.9, 0.35)
    t.paragraphs[0].alignment = PP_ALIGN.CENTER
    _run(t.paragraphs[0], txt, 13, col, bold=True)
t = _box(s, 0.7, 5.85, 11.9, 0.9)
t.paragraphs[0].alignment = PP_ALIGN.CENTER
_run(t.paragraphs[0], "A cycle, not a line — specialized roles · separate Nemotron calls · shared H-MEM evidence "
                      "memory · full trace + signed Meterless receipts.", 13, MUTED)
_footer(s)

# 6 — Cross-source (two columns)
s = prs.slides.add_slide(BLANK)
_bg(s, WHITE)
_header(s, "The innovation", "The reconciliation is the insight")
col_w, cy = 5.6, 2.45
for cx, head, items, col in [
    (0.7, "Trialed — ClinicalTrials.gov",
     ["tDCS · theta-burst · deep-brain stimulation", "24 device trial-arms for anorexia nervosa"], TEAL),
    (7.0, "Cleared — openFDA 510(k)",
     ["25 cleared transcranial neuro-devices", "specialties: Neurology · Cardiology · Radiology",
      "none reference anorexia"], GOLD),
]:
    _rect(s, cx, cy, col_w, 0.06, col)
    tf = _box(s, cx, cy + 0.15, col_w, 0.5)
    _run(tf.paragraphs[0], head, 16, INK, bold=True)
    tf = _box(s, cx, cy + 0.8, col_w, 2.6)
    f = True
    for it in items:
        p = tf.paragraphs[0] if f else tf.add_paragraph()
        f = False
        p.space_after = Pt(10)
        _run(p, "•  ", 16, col, bold=True)
        _run(p, it, 16, INK)
tf = _box(s, 0.7, 5.7, 11.9, 0.9)
_run(tf.paragraphs[0], "Trialed ≠ cleared.  ", 18, RED, bold=True)
_run(tf.paragraphs[0], "The agent surfaces the gap explicitly — that reconciliation is the answer.", 18, INK)
_footer(s)

# 7 — The integrity moment: the climax — dark, dramatic, the agent catching a false claim
s = prs.slides.add_slide(BLANK)
_bg(s, INK)
t = _box(s, 0.7, 0.55, 11.9, 0.4)
_run(t.paragraphs[0], "THE INTEGRITY MOMENT", 13, RED_BRIGHT, bold=True)
t = _box(s, 0.7, 1.0, 12.0, 1.0)
_run(t.paragraphs[0], "It deletes what it can't prove.", 36, WHITE, bold=True)
_rect(s, 0.72, 1.98, 1.1, 0.055, RED_BRIGHT)
t = _box(s, 0.7, 2.75, 11.0, 0.4)
_run(t.paragraphs[0], "The agent was handed a plausible, false claim:", 16, DIM)
t = _box(s, 0.7, 3.3, 8.7, 1.0)
_run(t.paragraphs[0], "rTMS is FDA-cleared to treat anorexia nervosa", 27, DIM, bold=True, strike=True)
_stamp(s, 9.75, 3.05, 2.75, 1.05, "✗  UNSUPPORTED", RED_BRIGHT, rot=-9, size=16)
t = _box(s, 0.7, 4.55, 11.4, 1.4)
_run(t.paragraphs[0], "The cited records show those devices are cleared — but not one references anorexia. "
                      "Unsupported, so the agent refuses to say it.", 18, RGBColor(0xDD, 0xE6, 0xEB))
t = _box(s, 0.7, 6.15, 11.9, 0.5)
_run(t.paragraphs[0], "Every rejection ships with its reason — that is the “why.”", 15, TEAL, bold=True)

# 8 — Stack & sponsor tech: logo strip (recognition) + how each is used
S8_BULLETS = [
    ("Zetaris", "— governed SQL data layer; read-only guard; a cross-source join of two disconnected registries."),
    ("Meterless", "— H-MEM evidence engine + Nightshift provider; every model call returns a signed receipt."),
    ("NVIDIA Nemotron", "— pinned reasoning model (nemotron-3.5-lightning-30b-a3b); it does the adjudication."),
    ("Cursor", "— development environment."),
    ("Telemetry per run:", "24 Meterless-routed calls · 41,480 tokens · signed receipts."),
]
if all(os.path.exists(p) for p in (L_ZETARIS, L_METERLESS, L_NVIDIA, L_CURSOR)):
    s = prs.slides.add_slide(BLANK)
    _bg(s, WHITE)
    _header(s, "Sponsor tech", "Built on the sponsor stack")
    _logo_cards(s, [L_ZETARIS, L_METERLESS, L_NVIDIA, L_CURSOR], top=2.15, cell_h=0.95)
    _bullets(s, S8_BULLETS, top=3.4, size=15, gap=10, height=3.3)
    _footer(s)
else:
    content_slide("Sponsor tech", "Built on the sponsor stack", S8_BULLETS)

# 9 — unique + verified, with a ✓ VERIFIED stamp (the receipt motif)
s = prs.slides.add_slide(BLANK)
_bg(s, WHITE)
_header(s, "What's unique — and verified · Technical", "Unique, and proven to work")
_bullets(s, [
    ("Claim-level citation binding.", "Negative / absence evidence is first-class."),
    ("Self-consistency gate.", "Approvals are re-checked by independent sampled votes."),
    ("H-MEM as evidence memory.", "Mandatory source/provenance; provable Meterless receipts."),
    ("Verified:", "ruff-clean · GitHub Actions CI · 15 offline + live integration tests · "
                  "eval scorecard 8/8 · SAMPLE_MODE replay."),
], top=2.15, height=4.7, anchor=MSO_ANCHOR.MIDDLE, width=9.1)
_stamp(s, 10.1, 2.5, 2.5, 1.0, "✓  VERIFIED", TEAL, rot=-8, size=16)
_footer(s)

# 10 — Close: dark bookend — the resonant button, with the vision folded in
s = prs.slides.add_slide(BLANK)
_bg(s, INK)
t = _box(s, 0.9, 0.8, 11.5, 0.5)
_run(t.paragraphs[0], "WHERE THIS GOES", 13, TEAL, bold=True)
_rect(s, 0.92, 2.35, 1.4, 0.06, TEAL)
t = _box(s, 0.9, 2.6, 11.6, 1.3)
_run(t.paragraphs[0], "An answer you can check.", 46, WHITE, bold=True)
t = _box(s, 0.9, 4.2, 11.6, 1.1)
_run(t.paragraphs[0], "Next: richer indication signals · more conditions & sources · a multi-judge panel — "
                      "trustworthy, auditable AI for clinical evidence.", 18, DIM)
t = _box(s, 0.9, 6.5, 11.6, 0.5)
_run(t.paragraphs[0], "Fylaktis Fylaktou · UCLA Semel Institute · Open Agent Hackathon 2026 · Track 02", 13, DIM)

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Evidence-Intelligence-Agent.pptx")
prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides._sldIdLst)} slides)")
